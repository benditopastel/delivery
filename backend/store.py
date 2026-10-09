"""Excel test adapter. Replace this class with Supabase implementation later."""
from pathlib import Path
from threading import RLock
from openpyxl import Workbook, load_workbook
from datetime import datetime
from uuid import uuid4

SHEETS = {
 'categories':['id','name','description','position','active','image','group_ids'],
 'products':['id','category_id','name','description','price','sale_price','badge','active','stock_enabled','stock','images','addons','position','group_ids'],
 'addons':['id','name','price','active','stock_enabled','stock'],
 'addon_groups':['id','name','required','min_select','max_select','active','position','options','free_label'],
 'orders':['id','customer','phone','delivery','address','payment','items','subtotal','fee','discount','total','status','created_at','change_for','change_amount','coupon'],
 'banners':['id','title','subtitle','image','active','target'],
 'coupons':['id','code','type','value','min_total','active','include_products','exclude_products','include_categories','exclude_categories','exclude_sale','max_discount','usage_limit','per_phone_limit','starts_at','ends_at'],
 'settings':['key','value'],
 'loyalty':['id','phone','points'],
 'rewards':['id','name','kind','product_id','value','points','active'],
}
SEED={
 'categories':[
 ['c1','Combos','Os favoritos da casa',1,True,''],['c2','Hambúrgueres','Sabor em cada mordida',2,True,''],['c3','Porções','Para compartilhar',3,True,''],['c4','Bebidas','Para acompanhar',4,True,'']],
 'products':[
 ['p1','c1','Combo Suprema','Hambúrguer artesanal, batata crocante e refrigerante.',39.9,34.9,'Mais vendido',True,False,0,'','a1,a2'],
 ['p2','c1','Combo Chicken','Frango empanado, batatas e molho especial.',36.9,0,'Destaque',True,False,0,'','a2'],
 ['p3','c2','Burger Clássico','Pão brioche, carne 160g, queijo, alface e tomate.',28.9,0,'',True,False,0,'','a1,a2,a3'],
 ['p4','c2','Double Bacon','Dois burgers, muito cheddar e bacon crocante.',42.9,38.9,'Novidade',True,False,0,'','a1,a2,a3'],
 ['p5','c3','Batata Especial','Batatas crocantes com ervas e sal.',21.9,0,'',True,False,0,'','a2,a3'],
 ['p6','c4','Refrigerante Lata','Lata 350 ml bem gelada.',7.9,0,'',True,True,50,'','']],
 'addons':[['a1','Bacon extra',5,True,False,0],['a2','Cheddar extra',4,True,False,0],['a3','Molho da casa',3,True,False,0]],
 'banners':[['b1','O sabor que conquista','Combos especiais para deixar seu dia melhor.','',True,'c1']],
 'coupons':[['cup1','BEMVINDO10','percent',10,30,True]],
 'settings':[['store_name','Brasa Burger'],['delivery_fee','6'],['free_shipping','75'],['open','true'],['notice','Feito na hora, com muito sabor!'],['primary','#961406'],['secondary','#d5411a'],['pix_key','benditopastelcm@gmail.com']],
}
class ExcelStore:
 def __init__(self,path):
  self.path=Path(path); self.lock=RLock(); self.path.parent.mkdir(parents=True,exist_ok=True)
  if not self.path.exists():
   wb=Workbook();wb.remove(wb.active)
   for name,cols in SHEETS.items():
    ws=wb.create_sheet(name);ws.append(cols)
    for row in SEED.get(name,[]):ws.append(row)
   wb.save(self.path)
  else:
   with self.lock:
    wb=load_workbook(self.path);changed=False
    for name,cols in SHEETS.items():
     if name not in wb.sheetnames:
      wb.create_sheet(name).append(cols);changed=True
     else:
      ws=wb[name];existing=[c.value for c in ws[1]]
      for col in cols:
       if col not in existing:ws.cell(1,ws.max_column+1,col);changed=True
    if changed:wb.save(self.path)
    wb.close()
 def all(self,table):
  with self.lock:
   wb=load_workbook(self.path,read_only=True,data_only=True)
   ws=wb[table]; rows=list(ws.values);wb.close()
   return [dict(zip(rows[0],r)) for r in rows[1:] if any(x is not None for x in r)]
 def insert(self,table,data):
  with self.lock:
   wb=load_workbook(self.path);cols=SHEETS[table]
   item={k:data.get(k,'') for k in cols}; item['id']=item.get('id') or uuid4().hex[:10]
   wb[table].append([item.get(k,'') for k in cols]);wb.save(self.path);wb.close();return item
 def update(self,table,ident,data):
  with self.lock:
   wb=load_workbook(self.path);ws=wb[table];cols=SHEETS[table];key='key' if table=='settings' else 'id'
   for row in ws.iter_rows(min_row=2):
    if str(row[cols.index(key)].value)==str(ident):
     for k,v in data.items():
      if k in cols and k!=key:row[cols.index(k)].value=v
     wb.save(self.path);wb.close();return True
   wb.close();return False
 def delete(self,table,ident):
  with self.lock:
   wb=load_workbook(self.path);ws=wb[table];key='key' if table=='settings' else 'id';idx=SHEETS[table].index(key)+1
   for i in range(2,ws.max_row+1):
    if str(ws.cell(i,idx).value)==str(ident):ws.delete_rows(i);wb.save(self.path);wb.close();return True
   wb.close();return False
 def settings(self):return {x['key']:x['value'] for x in self.all('settings')}
 def setting(self,key,value):
  if not self.update('settings',key,{'value':str(value)}):self.insert('settings',{'key':key,'value':str(value)})
